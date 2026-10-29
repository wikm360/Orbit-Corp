import docx
import pytest
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.features.documents.ingestion import docx_pagination as pag
from app.features.documents.ingestion.parser import parse_document_pages


def _add_rendered_break(paragraph):
    """What Word writes where it last broke a page: an empty marker element at the start of a run."""
    run = paragraph.runs[0]
    run._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))


def _doc(tmp_path, build):
    document = docx.Document()
    build(document)
    path = tmp_path / "d.docx"
    document.save(path)
    return str(path)


def test_word_rendered_markers_give_the_page_numbers_word_showed(tmp_path):
    def build(d):
        d.add_paragraph("alpha one")
        d.add_paragraph("alpha two")
        _add_rendered_break(d.add_paragraph("beta one"))
        d.add_paragraph("beta two")
        _add_rendered_break(d.add_paragraph("gamma one"))

    pages = parse_document_pages(_doc(tmp_path, build))

    assert [p.page_number for p in pages] == [1, 2, 3]
    assert pages[0].text == "alpha one\n\nalpha two" and pages[1].text == "beta one\n\nbeta two" and pages[2].text == "gamma one"
    assert pages[0].metadata == {"page": 1, "pagination": "word"}


def test_a_marker_inside_a_paragraph_splits_it_at_a_word_boundary(tmp_path):
    def build(d):
        p = d.add_paragraph()
        p.add_run("first half of the sentence ")
        second = p.add_run("second half here")
        second._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))

    pages = parse_document_pages(_doc(tmp_path, build))

    assert [p.text for p in pages] == ["first half of the sentence", "second half here"]


def test_explicit_break_plus_rendered_marker_is_one_break_not_an_empty_page(tmp_path):
    def build(d):
        d.add_paragraph("before").add_run().add_break(WD_BREAK.PAGE)
        _add_rendered_break(d.add_paragraph("after"))

    pages = parse_document_pages(_doc(tmp_path, build))

    assert [p.text for p in pages] == ["before", "after"]


def test_without_markers_text_is_spread_over_the_declared_page_count(tmp_path):
    def build(d):
        for i in range(40):
            d.add_paragraph(f"paragraph {i} " + "word " * 80)

    path = _doc(tmp_path, build)
    estimated = parse_document_pages(path)
    assert all(p.metadata["pagination"] == "estimated" for p in estimated)
    assert len(estimated) == 10  # ~16.4k chars / 1800 per page; python-docx's stamped 'Pages=1' is implausible, so ignored

    from app.features.documents.ingestion.parser import _parse_docx_pages

    hinted = _parse_docx_pages(path, page_hint=12)
    assert len(hinted) == 12 and sum(p.text.count("paragraph") for p in hinted) == 40  # nothing lost or duplicated


def test_short_document_stays_one_page(tmp_path):
    pages = parse_document_pages(_doc(tmp_path, lambda d: d.add_paragraph("short")))
    assert len(pages) == 1 and pages[0].text == "short"


def test_table_rows_follow_the_page_they_start_on(tmp_path):
    def build(d):
        d.add_paragraph("intro")
        table = d.add_table(rows=3, cols=2)
        for i, row in enumerate(table.rows):
            row.cells[0].text, row.cells[1].text = f"r{i}a", f"r{i}b"
        table.rows[2].cells[0].paragraphs[0].runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))

    pages = parse_document_pages(_doc(tmp_path, build))

    assert "r0a | r0b" in pages[0].text and "r1a | r1b" in pages[0].text
    assert pages[1].text == "r2a | r2b"


def test_declared_page_counts(tmp_path):
    rtf = tmp_path / "a.rtf"
    rtf.write_bytes(rb"{\rtf1{\info{\nofpages41}}text}")
    assert pag.legacy_declared_pages(str(rtf), ".rtf") == 41
    assert pag.legacy_declared_pages(str(tmp_path / "missing.doc"), ".doc") is None
    assert pag.docx_declared_pages(str(tmp_path / "missing.docx")) is None
