"""Structured, queryable view of spreadsheet documents (xlsx / csv / legacy Excel) for the agent's sheet tools.

Why not just read the extracted text: a wide sheet flattened to `a | b | | | c` forces the model to count pipes and tell
two near-identical numbers apart, and it answered with another row's value. Here each sheet becomes a real table with
named columns, SQL runs on it (aggregation done by the database, not in the model's head), and every row keeps its
original row number so an answer can cite `Sheet1, row 12, column F`.

Pieces: `load_tables` (parse + sidecar cache), `normalize_text` (Persian letter/digit normalization, applied to stored
values and to the string literals of a query), `run_query` (read-only SQLite sandbox).
"""

import difflib
import json
import logging
import re
import sqlite3
import time
from dataclasses import dataclass, field
from pathlib import Path

import sqlglot
from sqlglot import exp
from sqlglot.errors import SqlglotError

from app.core.exceptions import BadRequestError

logger = logging.getLogger(__name__)

TABLES_VERSION = 1
MAX_ROWS_PER_SHEET = 20000
MAX_SQL_LENGTH = 4000
_HEADER_SCAN_ROWS = 10
_FUZZY_MAX_DISTINCT = 5000

_FA_TRANSLATE = str.maketrans(
    {
        "ي": "ی",
        "ك": "ک",
        "ى": "ی",
        "ۀ": "ه",
        "‌": " ",
        " ": " ",
        **{chr(0x06F0 + i): str(i) for i in range(10)},  # ۰-۹
        **{chr(0x0660 + i): str(i) for i in range(10)},  # ٠-٩
        "٬": ",",
        "٫": ".",
    }
)
_SPACES = re.compile(r"\s+")
_NUMERIC_TEXT = re.compile(r"^[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?$|^[+-]?\d+(?:\.\d+)?$")


def normalize_text(value: str) -> str:
    """Arabic/Persian letter variants, Persian/Arabic digits, ZWNJ and runs of whitespace -> one canonical form."""
    return _SPACES.sub(" ", value.translate(_FA_TRANSLATE)).strip()


def column_letter(index: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA."""
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


@dataclass
class Column:
    name: str
    letter: str
    type: str  # "INTEGER" | "REAL" | "TEXT"
    synthetic: bool = False  # name invented because the header cell was empty or only a number


@dataclass
class Table:
    name: str
    header_row: int  # 1-based row number of the header in the file
    columns: list[Column]
    rows: list[list]  # values aligned with `columns`
    row_numbers: list[int]  # original 1-based row number of each row


@dataclass
class QueryResult:
    columns: list[str]
    rows: list[list]
    truncated: bool
    suggestions: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _to_number(value):
    """int/float for numeric cells and numeric-looking text ('15,979,160', '۱۲٫۵'); None otherwise (leading zeros stay text)."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, str):
        text = normalize_text(value)
        if not _NUMERIC_TEXT.match(text):
            return None
        bare = text.lstrip("+-")
        if len(bare) > 1 and bare.startswith("0") and not bare.startswith("0."):
            return None
        text = text.replace(",", "")
        try:
            number = float(text)
        except ValueError:
            return None
        return int(number) if number.is_integer() and "." not in text else number
    return None


def _clean_cell(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if hasattr(value, "isoformat"):
        iso = value.isoformat()
        return iso[:-9] if iso.endswith("T00:00:00") else iso.replace("T", " ")
    text = normalize_text(str(value))
    return text or None


def _pick_header(raw: list[tuple[int, list]]) -> int:
    """Index (into `raw`) of the header row: the first of the first few rows with 2+ filled cells (skips title lines)."""
    for position, (_, cells) in enumerate(raw[:_HEADER_SCAN_ROWS]):
        if sum(1 for c in cells if c is not None) >= 2:
            return position
    return 0


def _build_table(name: str, raw: list[tuple[int, list]]) -> Table | None:
    raw = [(n, cells) for n, cells in raw if any(c is not None for c in cells)]
    if not raw:
        return None
    header_pos = _pick_header(raw)
    header_number, header_cells = raw[header_pos]
    body = raw[header_pos + 1 :]
    width = max([len(header_cells)] + [len(cells) for _, cells in body])

    names: list[str] = []
    synthetic: list[bool] = []
    seen: set[str] = set()
    for index in range(width):
        label = header_cells[index] if index < len(header_cells) else None
        letter = column_letter(index)
        is_synthetic = False
        if label is None:
            column_name = f"ستون_{letter}"
            is_synthetic = True
        elif isinstance(label, (int, float)):
            # A bare number as a header (e.g. a day count) says nothing about the column on its own.
            column_name = f"ستون_{letter} ({_clean_cell(label)})"
            is_synthetic = True
        else:
            column_name = str(label).replace('"', "'")
        base, suffix = column_name, 2
        while column_name.lower() in seen:
            column_name = f"{base} ({suffix})"
            suffix += 1
        seen.add(column_name.lower())
        names.append(column_name)
        synthetic.append(is_synthetic)

    columns: list[Column] = []
    rows: list[list] = [[None] * width for _ in body]
    for index in range(width):
        values = [cells[index] if index < len(cells) else None for _, cells in body]
        present = [v for v in values if v is not None]
        numbers = [_to_number(v) for v in present]
        if present and all(n is not None for n in numbers):
            kind = "REAL" if any(isinstance(n, float) for n in numbers) else "INTEGER"
            converted = iter(numbers)
            values = [next(converted) if v is not None else None for v in values]
        else:
            kind = "TEXT"
            values = [v if v is None else str(_clean_cell(v)) for v in values]
        for row, value in zip(rows, values):
            row[index] = value
        columns.append(Column(name=names[index], letter=column_letter(index), type=kind, synthetic=synthetic[index]))

    return Table(name=name, header_row=header_number, columns=columns, rows=rows, row_numbers=[n for n, _ in body])


def read_xlsx_tables(file_path: str) -> list[Table]:
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, data_only=True, read_only=True)
    tables: list[Table] = []
    try:
        for sheet in workbook.worksheets:
            if getattr(sheet, "reset_dimensions", None):
                sheet.reset_dimensions()
            raw: list[tuple[int, list]] = []
            for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [_clean_cell(v) for v in values]
                if any(c is not None for c in cells):
                    raw.append((number, cells))
                if len(raw) >= MAX_ROWS_PER_SHEET:
                    logger.warning("Sheet %r of %s truncated at %d rows", sheet.title, file_path, MAX_ROWS_PER_SHEET)
                    break
            table = _build_table(sheet.title, raw)
            if table is not None:
                tables.append(table)
    finally:
        workbook.close()
    return tables


def read_csv_tables(file_path: str) -> list[Table]:
    import csv
    import io

    from app.features.documents.ingestion.parser import _parse_txt

    text = _parse_txt(file_path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    raw: list[tuple[int, list]] = []
    for number, values in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        cells = [_clean_cell(v) for v in values]
        if any(c is not None for c in cells):
            raw.append((number, cells))
        if len(raw) >= MAX_ROWS_PER_SHEET:
            break
    table = _build_table(Path(file_path).stem, raw)
    return [table] if table is not None else []


def _read_tables(file_path: str) -> list[Table]:
    from app.features.documents.ingestion.parser import LEGACY_OFFICE_EXTENSIONS

    suffix = Path(file_path).suffix.lower()
    if suffix == ".xlsx":
        return read_xlsx_tables(file_path)
    if suffix == ".csv":
        return read_csv_tables(file_path)
    if suffix in LEGACY_OFFICE_EXTENSIONS and suffix in {".xls", ".xlsm", ".ods"}:
        from app.features.documents.ingestion.converter import converted_copy

        with converted_copy(file_path) as modern:
            return read_xlsx_tables(modern)
    raise BadRequestError("This document is not a spreadsheet (xlsx / xls / csv)")


def is_spreadsheet(file_path: str) -> bool:
    return Path(file_path).suffix.lower() in {".xlsx", ".xls", ".xlsm", ".ods", ".csv"}


# ---------------------------------------------------------------------------
# Sidecar cache
# ---------------------------------------------------------------------------


def tables_cache_path(file_path: str) -> Path:
    return Path(f"{file_path}.tables.json")


def cleanup_tables_cache(file_path: str) -> None:
    try:
        path = tables_cache_path(file_path)
        if path.exists():
            path.unlink()
    except Exception as exc:
        logger.debug("Failed deleting tables cache for %s: %s", file_path, exc)


def _to_json(tables: list[Table], source_size: int) -> dict:
    return {
        "version": TABLES_VERSION,
        "source_size": source_size,
        "tables": [
            {
                "name": t.name,
                "header_row": t.header_row,
                "columns": [vars(c) for c in t.columns],
                "rows": t.rows,
                "row_numbers": t.row_numbers,
            }
            for t in tables
        ],
    }


def _from_json(data: dict) -> list[Table]:
    return [
        Table(
            name=t["name"],
            header_row=t["header_row"],
            columns=[Column(**c) for c in t["columns"]],
            rows=t["rows"],
            row_numbers=t["row_numbers"],
        )
        for t in data["tables"]
    ]


def load_tables(file_path: str) -> list[Table]:
    """Parsed tables of a spreadsheet file, built once and cached next to it (blocking: call from a thread)."""
    path = Path(file_path)
    size = path.stat().st_size
    cache = tables_cache_path(file_path)
    if cache.exists():
        try:
            data = json.loads(cache.read_text(encoding="utf-8"))
            if data.get("version") == TABLES_VERSION and data.get("source_size") == size:
                return _from_json(data)
        except Exception as exc:
            logger.warning("Ignoring unreadable tables cache %s: %s", cache, exc)

    tables = _read_tables(file_path)
    try:
        tmp = cache.with_suffix(cache.suffix + ".tmp")
        tmp.write_text(json.dumps(_to_json(tables, size), ensure_ascii=False), encoding="utf-8")
        tmp.replace(cache)
    except Exception as exc:
        logger.warning("Could not write tables cache %s: %s", cache, exc)
    return tables


# ---------------------------------------------------------------------------
# Text for ingestion (search / read tools)
# ---------------------------------------------------------------------------


def record_pages(table: Table, rows_per_page: int) -> list[tuple[str, str]]:
    """(text, "first-last file rows") per page. One line per row as `name: value | ...` with empty cells left out, so a value
    always travels with its column name instead of depending on its position among many separators."""
    pages = []
    total = len(table.rows)
    for lo in range(0, total, rows_per_page):
        hi = min(lo + rows_per_page, total)
        first, last = table.row_numbers[lo], table.row_numbers[hi - 1]
        lines = [
            f"=== Sheet: {table.name} (data rows {lo + 1}-{hi} of {total}; file rows {first}-{last}) ===",
            "(For an exact value, lookup, count or sum over this file use the query_spreadsheet tool.)",
        ]
        for number, row in zip(table.row_numbers[lo:hi], table.rows[lo:hi]):
            cells = [f"{column.name}: {value}" for column, value in zip(table.columns, row) if value is not None]
            lines.append(f"ردیف {number}: " + " | ".join(cells))
        pages.append(("\n".join(lines), f"{first}-{last}"))
    return pages


# ---------------------------------------------------------------------------
# Description (schema + samples) for the model
# ---------------------------------------------------------------------------


def _md_cell(value) -> str:
    return "" if value is None else str(value).replace("|", "/").replace("\n", " ")


def markdown_table(headers: list[str], rows: list[list]) -> str:
    lines = ["| " + " | ".join(_md_cell(h) for h in headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    lines.extend("| " + " | ".join(_md_cell(v) for v in row) + " |" for row in rows)
    return "\n".join(lines)


def describe_tables(tables: list[Table], sheet: str | None = None, sample_rows: int = 5) -> str:
    chosen = [t for t in tables if sheet is None or normalize_text(t.name).lower() == normalize_text(sheet).lower()]
    if not chosen:
        names = ", ".join(f'"{t.name}"' for t in tables)
        raise BadRequestError(f"Sheet not found. Available sheets: {names}")

    parts = [
        "Every sheet is a SQL table (use the sheet name in double quotes). Each table has a leading `_row` column = the "
        "row number in the original file; columns are named from the header row. Text is normalized (ي→ی, ك→ک, "
        "Persian digits→0-9, half-space→space)."
    ]
    for table in chosen:
        parts.append(f'\n### Sheet "{table.name}" — {len(table.rows)} data rows (header at file row {table.header_row})')
        info_rows = []
        for index, column in enumerate(table.columns):
            values = [row[index] for row in table.rows if row[index] is not None]
            if column.type == "TEXT":
                distinct = sorted({str(v) for v in values})
                detail = ("values: " + ", ".join(distinct[:12]) + (" …" if len(distinct) > 12 else "")) if 0 < len(distinct) <= 12 else f"{len(distinct)} distinct"
            else:
                detail = f"min {min(values)} … max {max(values)}" if values else ""
            label = column.name + (" *" if column.synthetic else "")
            info_rows.append([column.letter, label, column.type, f"{len(values)}/{len(table.rows)}", detail])
        parts.append(markdown_table(["file column", "name", "type", "filled", "summary"], info_rows))
        if any(c.synthetic for c in table.columns):
            parts.append("* = the header cell was empty or only a number; the name is synthetic, look at the sample rows to see what it holds.")
        sample = [[n] + row for n, row in zip(table.row_numbers[:sample_rows], table.rows[:sample_rows])]
        parts.append("First rows:\n" + markdown_table(["_row"] + [c.name for c in table.columns], sample))
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# SQL sandbox
# ---------------------------------------------------------------------------

_ALLOWED_FUNCTIONS = frozenset(
    {
        "abs", "round", "upper", "lower", "length", "substr", "substring", "trim", "ltrim", "rtrim", "replace", "instr",
        "coalesce", "ifnull", "nullif", "iif", "typeof", "min", "max", "sum", "avg", "count", "total", "group_concat",
        "row_number", "rank", "dense_rank", "ntile", "lag", "lead", "first_value", "last_value", "nth_value",
        "percent_rank", "cume_dist", "like", "glob", "printf", "format", "cast", "date", "strftime", "likely",
        "unlikely", "sign", "hex", "quote", "concat",
    }
)


def _authorizer(action, arg1, arg2, dbname, source):
    if action in (sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ):
        return sqlite3.SQLITE_OK
    if action == sqlite3.SQLITE_FUNCTION:
        return sqlite3.SQLITE_OK if (arg2 or "").lower() in _ALLOWED_FUNCTIONS else sqlite3.SQLITE_DENY
    return sqlite3.SQLITE_DENY


def _quote(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def _prepare_sql(sql: str) -> tuple[str, list[str]]:
    """Validates a single SELECT and normalizes its string literals; returns (sql, literals for 'did you mean')."""
    sql = sql.strip().rstrip(";").strip()
    if not sql:
        raise BadRequestError("'sql' is empty")
    if len(sql) > MAX_SQL_LENGTH:
        raise BadRequestError(f"SQL is too long (max {MAX_SQL_LENGTH} characters)")
    try:
        statements = sqlglot.parse(sql, read="sqlite")
    except SqlglotError as exc:
        raise BadRequestError(f"SQL syntax error: {exc}") from exc
    if len(statements) != 1 or statements[0] is None:
        raise BadRequestError("Exactly one SELECT statement is allowed")
    tree = statements[0]
    if not isinstance(tree, (exp.Select, exp.Union, exp.Subquery)):
        raise BadRequestError("Only SELECT statements are allowed")

    literals: list[str] = []
    for literal in tree.find_all(exp.Literal):
        if literal.is_string:
            normalized = normalize_text(literal.this)
            literal.set("this", normalized)
            bare = normalized.strip("%_ ")
            if len(bare) >= 3 and not bare.isdigit():
                literals.append(bare)
    return tree.sql(dialect="sqlite"), literals


def _suggest(tables: list[Table], literals: list[str], limit: int = 5) -> list[str]:
    found: list[str] = []
    for literal in literals:
        scored: list[tuple[float, str]] = []
        for table in tables:
            for index, column in enumerate(table.columns):
                if column.type != "TEXT":
                    continue
                distinct = {row[index] for row in table.rows if row[index]}
                if len(distinct) > _FUZZY_MAX_DISTINCT:
                    continue
                for candidate in difflib.get_close_matches(literal, distinct, n=limit, cutoff=0.7):
                    ratio = difflib.SequenceMatcher(None, literal, candidate).ratio()
                    scored.append((ratio, f'"{table.name}"."{column.name}": {candidate}'))
        for _, text in sorted(scored, reverse=True)[:limit]:
            if text not in found:
                found.append(text)
    return found[:limit]


def run_query(tables: list[Table], sql: str, *, max_rows: int, timeout_seconds: float) -> QueryResult:
    """Runs one read-only SELECT over the tables (blocking: call from a thread)."""
    prepared, literals = _prepare_sql(sql)

    conn = sqlite3.connect(":memory:")
    try:
        for table in tables:
            column_defs = ", ".join(f"{_quote(c.name)} {c.type}" for c in table.columns)
            conn.execute(f'CREATE TABLE {_quote(table.name)} ("_row" INTEGER, {column_defs})')
            placeholders = ", ".join("?" for _ in range(len(table.columns) + 1))
            conn.executemany(
                f"INSERT INTO {_quote(table.name)} VALUES ({placeholders})",
                ([number] + row for number, row in zip(table.row_numbers, table.rows)),
            )
        conn.commit()
        conn.execute("PRAGMA query_only = ON")
        conn.set_authorizer(_authorizer)
        deadline = time.monotonic() + timeout_seconds
        conn.set_progress_handler(lambda: 1 if time.monotonic() > deadline else 0, 10000)

        try:
            cursor = conn.execute(prepared)
            fetched = cursor.fetchmany(max_rows + 1)
        except sqlite3.OperationalError as exc:
            if "interrupted" in str(exc).lower():
                raise BadRequestError(f"Query exceeded the {timeout_seconds:g}s limit; simplify it") from exc
            columns_hint = "; ".join(
                f'"{t.name}"({", ".join(_quote(c.name) for c in t.columns)})' for t in tables
            )
            raise BadRequestError(f"SQL error: {exc}. Tables and columns: {columns_hint}") from exc
        except sqlite3.DatabaseError as exc:
            raise BadRequestError(f"SQL not allowed or failed: {exc}") from exc

        names = [d[0] for d in cursor.description or []]
        truncated = len(fetched) > max_rows
        rows = [list(r) for r in fetched[:max_rows]]
        suggestions = _suggest(tables, literals) if not rows and literals else []
        return QueryResult(columns=names, rows=rows, truncated=truncated, suggestions=suggestions)
    finally:
        conn.close()


def column_letters_for(tables: list[Table], names: list[str]) -> dict[str, str]:
    """Name -> file column letter for result columns that are plain columns of some sheet (for citations)."""
    letters: dict[str, str] = {}
    for table in tables:
        for column in table.columns:
            letters.setdefault(column.name, column.letter)
    return {name: letters[name] for name in names if name in letters}
