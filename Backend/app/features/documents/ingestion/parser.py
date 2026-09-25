from pathlib import Path

from app.core.exceptions import BadRequestError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".txt"}


def parse_document(file_path: str) -> str:
    """Extract plain text from a supported document format.

    Each format uses a pure-Python library (no system dependencies like
    poppler/tesseract), which matches the MVP scope of text-based files only
    (no OCR / scanned PDFs).
    """
    suffix = Path(file_path).suffix.lower()

    if suffix == ".pdf":
        return _parse_pdf(file_path)
    if suffix == ".docx":
        return _parse_docx(file_path)
    if suffix == ".pptx":
        return _parse_pptx(file_path)
    if suffix == ".xlsx":
        return _parse_xlsx(file_path)
    if suffix == ".txt":
        return _parse_txt(file_path)

    raise BadRequestError(f"Unsupported file type: {suffix}")


def _parse_txt(file_path: str) -> str:
    """Decodes a .txt file, detecting UTF-16 (common for Windows-originated
    files - e.g. PowerShell's `Start-Transcript` writes UTF-16LE) instead of
    assuming UTF-8 outright.

    A blind `read_text(encoding="utf-8")` doesn't raise on UTF-16 content:
    every ASCII byte in it is followed by a `\\x00` byte, and `\\x00` alone
    is *also* valid UTF-8 (it decodes to U+0000), so the read silently
    "succeeds" while leaving a literal NUL between every character. Postgres
    then rejects the whole chunk outright - `text`/`varchar` columns can
    never contain an embedded NUL byte, regardless of the source encoding.
    """
    raw = Path(file_path).read_bytes()

    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = raw.decode("utf-16")
    elif raw.startswith(b"\xef\xbb\xbf"):
        text = raw.decode("utf-8-sig")
    else:
        # No BOM: guess UTF-16LE by its tell - a NUL byte after most
        # characters. Genuine UTF-8 text (Latin or multi-byte, Persian
        # included) essentially never contains NUL bytes at all.
        sample = raw[:4000]
        if sample.count(b"\x00") > len(sample) // 4:
            try:
                text = raw.decode("utf-16-le")
            except UnicodeDecodeError:
                text = raw.decode("utf-8", errors="ignore")
        else:
            text = raw.decode("utf-8", errors="ignore")

    return text.replace("\x00", "")


def _parse_pdf(file_path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(file_path: str) -> str:
    import docx

    document = docx.Document(file_path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _parse_pptx(file_path: str) -> str:
    from pptx import Presentation

    presentation = Presentation(file_path)
    lines: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


def _parse_xlsx(file_path: str) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
