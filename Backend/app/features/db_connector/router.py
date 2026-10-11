import asyncio
import csv
import io
import uuid

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.dependencies import get_current_user_context, require_admin, UserContext
from app.core.config import get_settings
from app.core.exceptions import BadRequestError, ForbiddenError, NotFoundError
from app.features.db_connector import hints, service
from app.features.db_connector.executor import DbExecutionError, data_freshness, execute_query
from app.features.db_connector.models import DbQueryLog, QueryStatus
from app.features.db_connector.sql_validator import SqlValidationError, validate_sql
from app.features.db_connector.schemas import (
    DbQueryFeedbackRequest,
    DbQueryLogRead,
    DbQueryStats,
    IngestResponse,
)

router = APIRouter(prefix="/accounting-db", tags=["accounting-db"])


async def _get_owned_log(db: AsyncSession, query_id: uuid.UUID, context: UserContext) -> DbQueryLog:
    """A query log is visible to the user who asked it (and to admins) - never to other users."""
    log = (await db.execute(select(DbQueryLog).where(DbQueryLog.id == query_id))).scalar_one_or_none()
    if log is None:
        raise NotFoundError("Query log not found")
    if log.user_id != context.id and not context.is_admin:
        raise ForbiddenError()
    return log


@router.post("/queries/{query_id}/feedback", response_model=DbQueryLogRead)
async def submit_query_feedback(
    query_id: uuid.UUID,
    payload: DbQueryFeedbackRequest,
    context: UserContext = Depends(get_current_user_context),
    db: AsyncSession = Depends(get_db),
):
    """Submits user feedback (thumbs up +1, thumbs down -1) on an accounting query result you asked."""
    await _get_owned_log(db, query_id, context)
    return await service.record_feedback(
        db,
        query_id=query_id,
        feedback=payload.feedback,
        note=payload.note,
    )


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _safe_cell(value):
    """Spreadsheet formula injection guard: text from the database that starts with = + @ is stored as text, not as a formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "@", "\t", "\r"):
        return "'" + value
    return value


def _build_xlsx(columns: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "result"
    ws.sheet_view.rightToLeft = True
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163B48")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_safe_cell(v) for v in row])
    ws.freeze_panes = "A2"
    for idx, name in enumerate(columns, start=1):
        sample = [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r) and r[idx - 1] is not None]
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max([len(name)] + sample) + 3, 60)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@router.get("/queries/{query_id}/export")
async def export_query_result(
    query_id: uuid.UUID,
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    context: UserContext = Depends(get_current_user_context),
    db: AsyncSession = Depends(get_db),
):
    """Re-runs the stored (re-validated) SQL of one of your answers and returns up to `rahkaran_export_max_rows` rows as Excel/CSV."""
    log = await _get_owned_log(db, query_id, context)
    if not log.sql_text:
        raise BadRequestError("This answer has no SQL result to export")
    try:
        validated = validate_sql(log.sql_text)
        result = await execute_query(validated.sql, max_rows=get_settings().rahkaran_export_max_rows)
    except (SqlValidationError, DbExecutionError) as exc:
        raise BadRequestError(f"Export failed: {getattr(exc, 'message', str(exc))}")
    filename = f"accounting-result-{str(query_id)[:8]}"
    if format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(result.columns)
        writer.writerows([[_safe_cell(v) for v in row] for row in result.rows])
        body, media = ("\ufeff" + buffer.getvalue()).encode("utf-8"), "text/csv; charset=utf-8"
    else:
        body, media = await asyncio.to_thread(_build_xlsx, result.columns, result.rows), _XLSX
    return Response(
        content=body,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}.{format}"', "Cache-Control": "no-store"},
    )


@router.get("/queries", response_model=list[DbQueryLogRead], dependencies=[Depends(require_admin)])
async def list_queries(
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
    status: QueryStatus | None = None,
    db: AsyncSession = Depends(get_db),
):
    """Admin oversight: list historical accounting database queries and outcomes."""
    return await service.list_query_logs(db, limit=limit, offset=offset, status=status)


@router.get("/queries/stats", response_model=DbQueryStats, dependencies=[Depends(require_admin)])
async def get_stats(db: AsyncSession = Depends(get_db)):
    """Admin oversight: aggregate metrics on query volume, accuracy, and latency."""
    return await service.get_query_stats(db)


@router.post("/hints/ingest", response_model=IngestResponse, dependencies=[Depends(require_admin)])
async def trigger_ingest_hints(db: AsyncSession = Depends(get_db)):
    """Admin endpoint to re-sync and re-embed schema hints from the hints corpus."""
    stats = await hints.ingest_hints(db)
    return IngestResponse(**stats)


@router.get("/freshness")
async def get_freshness(_: UserContext = Depends(get_current_user_context)):
    """Returns snapshot freshness metadata for the Rahkaran database (authenticated users only)."""
    return await data_freshness()

